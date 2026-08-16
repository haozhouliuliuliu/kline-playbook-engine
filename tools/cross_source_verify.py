"""SHFE official cross-source verification for accepted canonical partitions.

For every accepted manifest partition with a comparable official SHFE daily
row, classifies MATCH / MISMATCH / NOT_COMPARABLE / NOT_AVAILABLE on
open/high/low/close/volume/open interest.  Reads the official annual raw
(2024 .xls / 2025 .xlsx) via xlrd/openpyxl (read-only reference parsing, same
as the fixture builder).  Conflicts are never averaged or silently repaired.

Usage (needs the reference-parsing venv with xlrd/openpyxl + tqsdk probe venv):
  python -m tools.cross_source_verify --manifest data/canonical/manifest.json
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path

import xlrd
import openpyxl

from src.cross_source_check import (
    DailyAggregate,
    OfficialDailyRow,
    aggregate_daily_from_bars,
    compare_daily,
    cross_source_summary,
)
from src.historical_manifest import DatasetManifest, ManifestPartition

PROJECT_ROOT = Path(__file__).resolve().parents[1]
REFERENCE_BASE = Path(r"C:\Users\H\Documents\steel-radar-lite\data\archive\gate1\raw\shfe-official")
ZIPS = {
    "2024": REFERENCE_BASE / "P020250411613533972715.zip",
    "2025": REFERENCE_BASE / "P020260116534160394483.zip",
}
CONTRACT_RE = re.compile(r"^([A-Za-z]{1,3})(\d{3,4})$")


def sheet_rows(member_bytes: bytes, member_name: str):
    lower = member_name.lower()
    if lower.endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(member_bytes), read_only=True, data_only=True)
        ws = wb["所内合约行情表"] if "所内合约行情表" in wb.sheetnames else wb.worksheets[0]
        for row in ws.iter_rows(min_row=5, values_only=True):
            yield row
    elif lower.endswith(".xls"):
        book = xlrd.open_workbook(file_contents=member_bytes)
        sh = book.sheet_by_index(0)
        for r in range(4, sh.nrows):
            yield [sh.cell_value(r, c) for c in range(sh.ncols)]
    else:
        return


def load_official_rows() -> dict[tuple[str, date], OfficialDailyRow]:
    """contract-code(lowercase) + date -> official daily row."""
    out: dict[tuple[str, date], OfficialDailyRow] = {}
    for year, zpath in ZIPS.items():
        with zipfile.ZipFile(zpath) as zf:
            for name in zf.namelist():
                if "ine" in name.lower():
                    continue
                if not name.lower().endswith((".xls", ".xlsx")):
                    continue
                if "所内合约行情报表" not in name and "所内合约行情表" not in name:
                    continue
                current = None
                for row in sheet_rows(zf.read(name), name):
                    if not row:
                        continue
                    if row[0] is not None and str(row[0]).strip():
                        current = str(row[0]).strip()
                    day_s = str(row[1]).strip() if row[1] is not None else ""
                    if not day_s.isdigit() or len(day_s) != 8:
                        continue
                    code = (current or "").lower()
                    m = CONTRACT_RE.match(code)
                    if not m:
                        continue
                    day = date(int(day_s[:4]), int(day_s[4:6]), int(day_s[6:8]))
                    try:
                        o = Decimal(str(row[4]))
                        h = Decimal(str(row[5]))
                        l = Decimal(str(row[6]))
                        c = Decimal(str(row[7]))
                        vol = int(float(row[11] or 0))
                        oi = Decimal(str(row[13])) if row[13] is not None and str(row[13]).strip() else None
                    except (TypeError, ValueError, ArithmeticError):
                        continue
                    out[(code, day)] = OfficialDailyRow(code, day, o, h, l, c, vol, oi)
    return out


def _num(value) -> Decimal | None:
    if value is None:
        return None
    return Decimal(str(value))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--raw-dir", default=str(PROJECT_ROOT / "data" / "raw" / "historical"))
    parser.add_argument("--out", default=None, help="write cross-source JSON summary")
    args = parser.parse_args()

    manifest = DatasetManifest.load(Path(args.manifest))
    official = load_official_rows()
    print(f"OFFICIAL_ROWS_LOADED={len(official)} manifest_partitions={len(manifest.partitions)}")

    from tools.build_canonical_dataset import _find_raw_file, SOURCE_SPEC  # noqa: PLC0415

    comparisons: list[dict] = []
    for partition in manifest.partitions:
        code = partition.contract.split(".")[1].lower()
        official_row = official.get((code, partition.trading_day))
        try:
            raw = _find_raw_file(Path(args.raw_dir), partition).read_bytes()
            # reconstruct a minimal bar view from the raw TqSdk JSON
            records = json.loads(raw.decode("utf-8"))
            bars = []
            for rec in records:
                bars.append(_BarView(rec))
            agg = aggregate_daily_from_bars(code, partition.trading_day, bars)
            comparisons.append(compare_daily(agg, official_row))
        except Exception as exc:  # noqa: BLE001
            comparisons.append({
                "contract": partition.contract, "day": partition.trading_day.isoformat(),
                "result": "NOT_COMPARABLE", "detail": f"{type(exc).__name__}: {exc}",
            })

    summary = cross_source_summary(comparisons)
    summary["partitions_checked"] = len(comparisons)
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    if args.out:
        Path(args.out).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out).write_text(
            json.dumps({"summary": summary, "comparisons": comparisons}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        print(f"WROTE {args.out}")


class _BarView:
    """Raw TqSdk record adapted to the aggregation surface."""

    def __init__(self, rec: dict) -> None:
        self.completed_bar = type("CB", (), {"start": _ts(rec["datetime"])})()
        self.open = _num(rec.get("open"))
        self.high = _num(rec.get("high"))
        self.low = _num(rec.get("low"))
        self.close = _num(rec.get("close"))
        self.volume = int(rec.get("volume") or 0)
        self.open_interest_end = _num(rec.get("close_oi"))


def _ts(ns: int):
    from datetime import datetime, timedelta  # noqa: PLC0415
    from zoneinfo import ZoneInfo  # noqa: PLC0415

    CN = ZoneInfo("Asia/Shanghai")
    return datetime.fromtimestamp(ns / 1_000_000_000, tz=CN) + timedelta(microseconds=(ns % 1_000_000_000) // 1000)


if __name__ == "__main__":
    main()
